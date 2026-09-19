import io
import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF byte content."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_pages.append(text.strip())
        return "\n\n".join(extracted_pages)
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX byte content."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from TXT or raw markdown."""
    for encoding in ["utf-8", "latin-1", "cp1252", "ascii"]:
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()

def parse_resume_content(filename: str, file_bytes: bytes) -> str:
    """Route file by extension and extract normalized text."""
    ext = Path(filename).suffix.lower()
    
    if ext == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        text = extract_text_from_docx(file_bytes)
    elif ext in [".txt", ".md", ".rtf"]:
        text = extract_text_from_txt(file_bytes)
    else:
        # Fallback attempt plain text
        try:
            text = extract_text_from_txt(file_bytes)
        except Exception:
            raise ValueError(f"Unsupported resume file format: {ext}. Please upload a PDF, DOCX, or TXT.")

    if not text or len(text.strip()) < 20:
        raise ValueError("The uploaded document appears to be empty or contains no readable text.")

    return text.strip()
