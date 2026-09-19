import os
import csv
from pathlib import Path
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.config import settings

HEADERS = [
    "Candidate ID",
    "Screened Date",
    "Full Name",
    "Email Address",
    "Phone Number",
    "Job Applied For",
    "Overall Score",
    "Screening Tier",
    "Skills Score",
    "Experience Score",
    "Matched Skills",
    "Missing Skills",
    "Top Strengths",
    "Gaps & Observations",
    "AI Executive Summary",
    "Notification Status"
]

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

SHORTLIST_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
SHORTLIST_FONT = Font(name="Calibri", size=10, bold=True, color="166534")

REVIEW_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
REVIEW_FONT = Font(name="Calibri", size=10, bold=True, color="92400E")

REJECT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
REJECT_FONT = Font(name="Calibri", size=10, bold=True, color="991B1B")

THIN_BORDER = Border(
    left=Side(style='thin', color="E2E8F0"),
    right=Side(style='thin', color="E2E8F0"),
    top=Side(style='thin', color="E2E8F0"),
    bottom=Side(style='thin', color="E2E8F0")
)

def init_or_load_workbook(filepath: str) -> openpyxl.Workbook:
    """Initialize workbook with styled header if it doesn't exist."""
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    
    if path.exists():
        try:
            wb = openpyxl.load_workbook(filepath)
            if "Screened Candidates" in wb.sheetnames:
                return wb
        except Exception:
            pass

    # Create fresh workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screened Candidates"
    ws.views.sheetView[0].showGridLines = True
    
    # Write and style headers
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 26
    wb.save(filepath)
    return wb

def sync_candidate_to_excel(candidate_record: Dict[str, Any], filepath: str = None) -> str:
    """Append or update a candidate screening row in the Excel workbook."""
    if not filepath:
        filepath = settings.EXCEL_PATH

    wb = init_or_load_workbook(filepath)
    ws = wb["Screened Candidates"]
    
    cid = candidate_record.get("id") or candidate_record.get("candidate_id")
    
    # Check if candidate row already exists
    target_row = None
    for r in range(2, ws.max_row + 1):
        cell_val = ws.cell(row=r, column=1).value
        if cell_val == cid:
            target_row = r
            break
            
    if not target_row:
        target_row = ws.max_row + 1 if ws.max_row >= 1 and ws.cell(row=1, column=1).value else 2
        
    matched_skills_str = ", ".join(candidate_record.get("matched_skills") or [])
    missing_skills_str = ", ".join(candidate_record.get("missing_skills") or [])
    strengths_str = " | ".join(candidate_record.get("strengths") or [])
    gaps_str = " | ".join(candidate_record.get("gaps_and_flags") or [])
    
    tier = candidate_record.get("recommendation", "Under Review")
    score = candidate_record.get("overall_score", 0.0)
    
    row_values = [
        cid,
        candidate_record.get("screened_at") or candidate_record.get("submitted_at") or "",
        candidate_record.get("full_name", ""),
        candidate_record.get("email", ""),
        candidate_record.get("phone", ""),
        candidate_record.get("job_title", ""),
        float(score),
        tier,
        float(candidate_record.get("skills_match_score", 0.0)),
        float(candidate_record.get("experience_score", 0.0)),
        matched_skills_str,
        missing_skills_str,
        strengths_str,
        gaps_str,
        candidate_record.get("summary_feedback", ""),
        candidate_record.get("notification_status", "Sent")
    ]
    
    for col_idx, val in enumerate(row_values, start=1):
        cell = ws.cell(row=target_row, column=col_idx, value=val)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")
        
        # Alignment & color formatting for specific columns
        if col_idx in [1, 7, 9, 10]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        if col_idx == 8: # Screening Tier
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if "shortlist" in tier.lower():
                cell.fill = SHORTLIST_FILL
                cell.font = SHORTLIST_FONT
            elif "not" in tier.lower() or "reject" in tier.lower():
                cell.fill = REJECT_FILL
                cell.font = REJECT_FONT
            else:
                cell.fill = REVIEW_FILL
                cell.font = REVIEW_FONT

    ws.row_dimensions[target_row].height = 22

    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(14, min(45, max_len + 3))

    wb.save(filepath)
    csv_path = str(Path(filepath).with_suffix(".csv"))
    sync_ws_to_csv(ws, csv_path)
    return filepath

def refresh_entire_excel(candidates: List[Dict[str, Any]], filepath: str = None) -> str:
    """Rebuild the Excel file cleanly from the latest candidate list."""
    if not filepath:
        filepath = settings.EXCEL_PATH

    # Delete existing or overwrite cleanly
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screened Candidates"
    ws.views.sheetView[0].showGridLines = True
    
    # Headers
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 26

    # Data rows
    for r_idx, c in enumerate(candidates, start=2):
        matched_skills_str = ", ".join(c.get("matched_skills") or [])
        missing_skills_str = ", ".join(c.get("missing_skills") or [])
        strengths_str = " | ".join(c.get("strengths") or [])
        gaps_str = " | ".join(c.get("gaps_and_flags") or [])
        tier = c.get("recommendation", "Under Review")
        score = c.get("overall_score", 0.0)

        row_values = [
            c.get("id"),
            c.get("screened_at") or c.get("submitted_at") or "",
            c.get("full_name", ""),
            c.get("email", ""),
            c.get("phone", ""),
            c.get("job_title", ""),
            float(score) if score is not None else 0.0,
            tier,
            float(c.get("skills_match_score", 0.0) or 0.0),
            float(c.get("experience_score", 0.0) or 0.0),
            matched_skills_str,
            missing_skills_str,
            strengths_str,
            gaps_str,
            c.get("summary_feedback", ""),
            c.get("notification_status", "Sent")
        ]
        
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if col_idx in [1, 7, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 8:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "shortlist" in str(tier).lower():
                    cell.fill = SHORTLIST_FILL
                    cell.font = SHORTLIST_FONT
                elif "not" in str(tier).lower() or "reject" in str(tier).lower():
                    cell.fill = REJECT_FILL
                    cell.font = REJECT_FONT
                else:
                    cell.fill = REVIEW_FILL
                    cell.font = REVIEW_FONT
        ws.row_dimensions[r_idx].height = 22

    # Column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(14, min(45, max_len + 3))

    wb.save(filepath)
    # Automatically export identical CSV file for immediate text/editor viewing
    csv_path = str(Path(filepath).with_suffix(".csv"))
    sync_ws_to_csv(ws, csv_path)
    return filepath

def sync_ws_to_csv(ws, csv_filepath: str) -> str:
    """Export worksheet rows directly to a clean CSV file."""
    Path(csv_filepath).parent.mkdir(parents=True, exist_ok=True)
    with open(csv_filepath, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([val if val is not None else "" for val in row])
    return csv_filepath
